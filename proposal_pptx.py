"""
BEKSCITY RUN CLUB — Executive Pitch Deck PowerPoint (.pptx) Generator
Uses the authentic master PowerPoint template:
'BCRC Partnership Proposal Run and Rumble.pptx'
Dynamically branding all 13 slides with target partner information,
custom logo embeds, event details, and co-founders' contacts.
"""

import os
import re
import io
import urllib.request
import pptx
from datetime import datetime
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))

# Candidate paths for master PPTX template
TEMPLATE_CANDIDATES = [
    os.path.join(HERE, "proposal_templates", "BCRC Partnership Proposal Run and Rumble.pptx"),
    os.path.expanduser("~/Downloads/BCRC Partnership Proposal Run and Rumble.pptx")
]


def find_pptx_template():
    for p in TEMPLATE_CANDIDATES:
        if os.path.exists(p) and os.path.getsize(p) > 1000000:
            return p
    return None


def fetch_image_bytes(img_source):
    """
    Loads raw image bytes from local disk path or remote URL.
    """
    if not img_source:
        return None

    try:
        if img_source.startswith("/static/"):
            local_path = os.path.join(HERE, img_source.lstrip("/"))
            if os.path.exists(local_path):
                with open(local_path, "rb") as f:
                    return f.read()

        if os.path.exists(img_source):
            with open(img_source, "rb") as f:
                return f.read()

        if img_source.startswith("http://") or img_source.startswith("https://"):
            # Normalize Wikimedia thumbnail width to 500px to avoid HTTP 400
            if "wikimedia.org" in img_source and "/thumb/" in img_source:
                img_source = re.sub(r"/\d+px-", "/500px-", img_source)
            req = urllib.request.Request(
                img_source,
                headers={"User-Agent": "BekscityRunClub/2.0 (contact@bekscity.run; runner-proposal-engine)"}
            )
            with urllib.request.urlopen(req, timeout=8) as res:
                return res.read()
    except Exception as e:
        print(f"[proposal_pptx] Warning fetching logo bytes from '{img_source}': {e}")

    return None


def build_proposal_pptx(output_path, proposal_data):
    """
    Loads master BCRC PowerPoint template and brands it for the target partner.
    """
    template_path = find_pptx_template()
    if not template_path:
        raise FileNotFoundError("Master PPTX template 'BCRC Partnership Proposal Run and Rumble.pptx' not found.")

    prs = pptx.Presentation(template_path)

    brand_name = (proposal_data.get("brand_name") or "PARTNER BRAND").strip()
    proposed_role = (proposal_data.get("proposed_role") or proposal_data.get("partnership_tier") or "Official Category Partner").strip()
    contact_person = (proposal_data.get("contact_person") or f"The {brand_name} Team").strip()
    event_name = (proposal_data.get("event_name") or "Run and Rumble").strip()
    
    # Format current date nicely e.g. "20th October 2026"
    now = datetime.now()
    day = now.day
    suffix = "th" if 11 <= day <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    date_str = f"{day}{suffix} {now.strftime('%B %Y')}"

    # Fetch logo bytes if available
    logo_bytes = fetch_image_bytes(proposal_data.get("brand_logo_url"))

    # Slide 1: Cover Page
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Prepared for:" in p.text:
                    p.text = f"Prepared for: {brand_name.upper()}  |   {date_str}"
        # Picture 7: Partner Brand Logo
        if shape.name == "Picture 7" and logo_bytes:
            try:
                img_part = slide1.part.related_part(shape._pic.blipFill.blip.rEmbed)
                img_part._blob = logo_bytes
            except Exception as e:
                print(f"[proposal_pptx] Slide 1 logo replace warning: {e}")

    # Slide 2: Founders Note & Proposal Summary Card
    if len(prs.slides) > 1:
        slide2 = prs.slides[1]
        for shape in slide2.shapes:
            if shape.has_text_frame:
                if shape.name == "Text 5" or shape.text_frame.text.strip() == "DOMINO’S PIZZA":
                    shape.text_frame.text = brand_name.upper()
                elif shape.name == "Text 7":
                    shape.text_frame.text = contact_person.upper()
                elif shape.name == "Text 9":
                    shape.text_frame.text = date_str

    # Slide 8: Upcoming Activation
    if len(prs.slides) > 7:
        slide8 = prs.slides[7]
        if event_name and event_name.lower() != "run and rumble" and event_name.lower() != "run & rumble":
            for shape in slide8.shapes:
                if shape.has_text_frame:
                    if shape.name == "Text 6":
                        shape.text_frame.text = event_name

    # Slide 13: Closing Contact Slide
    if len(prs.slides) > 12:
        slide13 = prs.slides[12]
        for shape in slide13.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if "07069582118" not in p.text and "PHONE" in p.text:
                        p.text = "07069582118, 08168120117, 08101674646"
            # Picture 16: Partner Brand Logo
            if shape.name == "Picture 16" and logo_bytes:
                try:
                    img_part13 = slide13.part.related_part(shape._pic.blipFill.blip.rEmbed)
                    img_part13._blob = logo_bytes
                except Exception as e:
                    print(f"[proposal_pptx] Slide 13 logo replace warning: {e}")

    # Global text replace of DOMINO'S across all remaining runs
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        for target in ["DOMINO’S PIZZA", "DOMINO'S PIZZA", "DOMINO’S", "DOMINO'S"]:
                            if target in r.text:
                                r.text = r.text.replace(target, brand_name)

    prs.save(output_path)
    return output_path
